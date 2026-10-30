"""
Minimalist & Clean PPTX Generator
Generates a modern 16:9 PowerPoint deck with clean typography and layout.
"""

import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE


def create_deck(title_text="Modern Presentation", output_path="presentation.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    COLOR_PRIMARY = RGBColor(0x11, 0x18, 0x27)
    COLOR_SECONDARY = RGBColor(0x6B, 0x72, 0x80)
    COLOR_ACCENT = RGBColor(0x25, 0x63, 0xEB)
    COLOR_CARD = RGBColor(0xF9, 0xFA, 0xFB)
    COLOR_BORDER = RGBColor(0xE5, 0xE7, 0xEB)

    blank_layout = prs.slide_layouts[6]

    # Slide 1: Cover
    slide1 = prs.slides.add_slide(blank_layout)
    title_box = slide1.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(10.5), Inches(3.0))
    tf = title_box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.font.name = "Arial"

    p2 = tf.add_paragraph()
    p2.text = "Crafted with Clean & Minimalist Aesthetics"
    p2.font.size = Pt(20)
    p2.font.color.rgb = COLOR_SECONDARY
    p2.font.name = "Arial"
    p2.space_before = Pt(14)

    # Slide 2: Two Column Content
    slide2 = prs.slides.add_slide(blank_layout)

    # Title
    t_box = slide2.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.0), Inches(1.0))
    tf2 = t_box.text_frame
    p_h = tf2.paragraphs[0]
    p_h.text = "Key Objectives & Focus"
    p_h.font.size = Pt(28)
    p_h.font.bold = True
    p_h.font.color.rgb = COLOR_PRIMARY
    p_h.font.name = "Arial"

    # Card 1 (Left)
    card1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.5))
    card1.fill.solid()
    card1.fill.fore_color.rgb = COLOR_CARD
    card1.line.color.rgb = COLOR_BORDER
    card1.line.width = Pt(1)

    tf_c1 = card1.text_frame
    tf_c1.word_wrap = True
    p_c1_t = tf_c1.paragraphs[0]
    p_c1_t.text = "01. Functional Minimalism"
    p_c1_t.font.size = Pt(20)
    p_c1_t.font.bold = True
    p_c1_t.font.color.rgb = COLOR_ACCENT

    p_c1_d = tf_c1.add_paragraph()
    p_c1_d.text = "\n• Eliminate unnecessary cognitive load\n• Structured visual whitespace\n• High-contrast legibility"
    p_c1_d.font.size = Pt(15)
    p_c1_d.font.color.rgb = COLOR_SECONDARY

    # Card 2 (Right)
    card2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.0), Inches(5.3), Inches(4.5))
    card2.fill.solid()
    card2.fill.fore_color.rgb = COLOR_CARD
    card2.line.color.rgb = COLOR_BORDER
    card2.line.width = Pt(1)

    tf_c2 = card2.text_frame
    tf_c2.word_wrap = True
    p_c2_t = tf_c2.paragraphs[0]
    p_c2_t.text = "02. Fast Automation"
    p_c2_t.font.size = Pt(20)
    p_c2_t.font.bold = True
    p_c2_t.font.color.rgb = COLOR_ACCENT

    p_c2_d = tf_c2.add_paragraph()
    p_c2_d.text = "\n• Programmatic slide assembly\n• Standardized layout tokens\n• Native editable PowerPoint output"
    p_c2_d.font.size = Pt(15)
    p_c2_d.font.color.rgb = COLOR_SECONDARY

    prs.save(output_path)
    print(f"Presentation generated successfully: {output_path}")


if __name__ == "__main__":
    out = "presentation_demo.pptx"
    if len(sys.argv) > 1:
        out = sys.argv[1]
    create_deck(output_path=out)
